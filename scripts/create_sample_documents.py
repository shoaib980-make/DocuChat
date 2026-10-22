
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)

sections = [
    (
        "Company Leave Policy",
        "Full-time employees receive 20 paid vacation days annually. "
        "Leave requests must be submitted through the HR portal and "
        "approved by the employee's manager.",
    ),
    (
        "Remote Work Policy",
        "Employees may request remote work with manager approval. "
        "Teams must maintain normal communication and availability "
        "during working hours.",
    ),
    (
        "Security Policy",
        "Employees must enable multi-factor authentication. "
        "Production credentials must be stored securely and must "
        "never be committed to Git.",
    ),
]

# Create DOCX
doc = Document()
doc.add_heading("Company Policies", level=1)

for heading, body in sections:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(body)

doc.save(DATA_DIR / "sample_company_policy.docx")

# Create PDF
pdf_path = DATA_DIR / "sample_company_policy.pdf"
pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
width, height = letter
y = height - 60

pdf.setFont("Helvetica-Bold", 16)
pdf.drawString(50, y, "Company Policies")
y -= 35

for heading, body in sections:
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, heading)
    y -= 22

    pdf.setFont("Helvetica", 10)

    # Wrap lines to keep text within the page width.
    words = body.split()
    line = ""

    for word in words:
        candidate = f"{line} {word}".strip()

        if pdf.stringWidth(candidate, "Helvetica", 10) > width - 100:
            pdf.drawString(50, y, line)
            y -= 15
            line = word
        else:
            line = candidate

    if line:
        pdf.drawString(50, y, line)
        y -= 28

pdf.save()

print("Created sample_company_policy.docx and sample_company_policy.pdf")
